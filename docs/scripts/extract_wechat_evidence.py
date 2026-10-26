"""Extract text from WeChat evidence files into plain text."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

base = Path(r"e:\ENT207TC\evidence\wechat-2026-09-10")
out = base / "_extracted"
out.mkdir(exist_ok=True)

try:
    from docx import Document

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    print("python-docx missing — falling back to XML parse")

try:
    import openpyxl

    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False
    print("openpyxl missing")


def docx_via_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab[^/]*/>", "\t", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_docx(path: Path) -> str:
    if not HAS_DOCX:
        return docx_via_xml(path)
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for ti, table in enumerate(doc.tables):
        parts.append(f"\n[TABLE {ti + 1}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


# unzip wireframes
zf = base / "UX_Design_Mobile_Low_Fidelity_Wireframes-English.zip"
if zf.exists():
    with zipfile.ZipFile(zf, "r") as z:
        z.extractall(base / "wireframes")
    names = sorted(p.name for p in (base / "wireframes").rglob("*") if p.is_file())
    print("unzipped wireframes:", len(names))
    for n in names[:40]:
        print("  ", n)

for p in sorted(base.glob("*.docx")):
    try:
        text = extract_docx(p)
        dest = out / f"{p.stem}.txt"
        dest.write_text(text, encoding="utf-8")
        print(f"OK docx {p.name} -> {len(text)} chars")
    except Exception as e:  # noqa: BLE001
        print(f"FAIL {p.name}: {e}")

if HAS_XLSX:
    for p in sorted(base.glob("*.xlsx")):
        wb = openpyxl.load_workbook(p, data_only=True)
        lines: list[str] = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            lines.append(f"===== SHEET: {sheet} =====")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in row):
                    lines.append("\t".join("" if c is None else str(c) for c in row))
        (out / f"{p.stem}.txt").write_text("\n".join(lines), encoding="utf-8")
        print(f"OK xlsx {p.name}")
else:
    try:
        import pandas as pd

        for p in sorted(base.glob("*.xlsx")):
            xls = pd.ExcelFile(p)
            lines = []
            for sheet in xls.sheet_names:
                df = pd.read_excel(p, sheet_name=sheet, header=None)
                lines.append(f"===== SHEET: {sheet} =====")
                lines.append(df.to_string())
            (out / f"{p.stem}.txt").write_text("\n".join(lines), encoding="utf-8")
            print(f"OK xlsx via pandas {p.name}")
    except Exception as e:  # noqa: BLE001
        print("xlsx fail", e)

print("extracted files:")
for p in sorted(out.glob("*.txt")):
    print(" -", p.name)
